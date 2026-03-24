#!/usr/bin/env python3
"""生成自研 PPTX 模板库。"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


TEMPLATE_SPECS = [
    {
        "name": "minimal-business",
        "title": "Minimal Business",
        "subtitle": "简洁商务风格",
        "palette": {
            "bg": RGBColor(250, 250, 252),
            "title": RGBColor(26, 26, 26),
            "body": RGBColor(60, 60, 60),
            "accent": RGBColor(28, 99, 235),
        },
    },
    {
        "name": "dark-contrast",
        "title": "Dark Contrast",
        "subtitle": "深色对比风格",
        "palette": {
            "bg": RGBColor(18, 22, 33),
            "title": RGBColor(255, 255, 255),
            "body": RGBColor(210, 216, 226),
            "accent": RGBColor(0, 194, 255),
        },
    },
    {
        "name": "warm-light",
        "title": "Warm Light",
        "subtitle": "温暖浅色风格",
        "palette": {
            "bg": RGBColor(255, 250, 243),
            "title": RGBColor(50, 43, 40),
            "body": RGBColor(96, 83, 77),
            "accent": RGBColor(230, 126, 34),
        },
    },
    {
        "name": "tech-neon",
        "title": "Tech Neon",
        "subtitle": "科技霓虹风格",
        "palette": {
            "bg": RGBColor(10, 12, 20),
            "title": RGBColor(233, 243, 255),
            "body": RGBColor(167, 184, 204),
            "accent": RGBColor(91, 255, 223),
        },
    },
    {
        "name": "elegant-serif",
        "title": "Elegant Serif",
        "subtitle": "典雅质感风格",
        "palette": {
            "bg": RGBColor(248, 246, 242),
            "title": RGBColor(36, 34, 32),
            "body": RGBColor(90, 85, 80),
            "accent": RGBColor(163, 122, 74),
        },
    },
    {
        "name": "fresh-green",
        "title": "Fresh Green",
        "subtitle": "清新自然风格",
        "palette": {
            "bg": RGBColor(246, 251, 248),
            "title": RGBColor(30, 51, 40),
            "body": RGBColor(76, 96, 86),
            "accent": RGBColor(46, 177, 120),
        },
    },
    {
        "name": "bold-corporate",
        "title": "Bold Corporate",
        "subtitle": "强对比企业风格",
        "palette": {
            "bg": RGBColor(244, 247, 252),
            "title": RGBColor(20, 32, 48),
            "body": RGBColor(69, 79, 92),
            "accent": RGBColor(255, 111, 0),
        },
    },
]


def _set_slide_background(slide, color: RGBColor) -> None:
    """设置幻灯片背景色。

    :param slide: 幻灯片对象
    :param color: 背景色
    :return: None
    """
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _style_title_shape(shape, color: RGBColor, size_pt: int) -> None:
    """设置标题文本样式。

    :param shape: 标题形状
    :param color: 字体颜色
    :param size_pt: 字号
    :return: None
    """
    if not shape:
        return
    if not shape.text_frame:
        return
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = color
            run.font.size = Pt(size_pt)


def _add_accent_bar(slide, color: RGBColor, y_inch: float, height_inch: float) -> None:
    """添加强调色条。

    :param slide: 幻灯片对象
    :param color: 色条颜色
    :param y_inch: 顶部位置（英寸）
    :param height_inch: 高度（英寸）
    :return: None
    """
    left = Inches(0)
    top = Inches(y_inch)
    width = Inches(13.33)
    height = Inches(height_inch)
    shape = slide.shapes.add_shape(1, left, top, width, height)
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = color
    shape.line.fill.background()


def _add_image_left_layout(slide, palette: dict[str, RGBColor]) -> None:
    """添加图左文右版式示例。

    :param slide: 幻灯片对象
    :param palette: 颜色方案
    :return: None
    """
    image_box = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.6), Inches(5.2), Inches(4.5)
    )
    image_box.fill.solid()
    image_box.fill.fore_color.rgb = palette["accent"]
    image_box.line.fill.background()

    text_box = slide.shapes.add_textbox(Inches(6.2), Inches(1.6), Inches(6.5), Inches(4.5))
    text_frame = text_box.text_frame
    text_frame.text = "要点一"
    text_frame.add_paragraph().text = "要点二"


def _add_image_right_layout(slide, palette: dict[str, RGBColor]) -> None:
    """添加图右文左版式示例。

    :param slide: 幻灯片对象
    :param palette: 颜色方案
    :return: None
    """
    text_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.6), Inches(6.2), Inches(4.5))
    text_frame = text_box.text_frame
    text_frame.text = "要点一"
    text_frame.add_paragraph().text = "要点二"

    image_box = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(7.1), Inches(1.6), Inches(5.2), Inches(4.5)
    )
    image_box.fill.solid()
    image_box.fill.fore_color.rgb = palette["accent"]
    image_box.line.fill.background()


def _add_chart_layout(slide, palette: dict[str, RGBColor]) -> None:
    """添加图表页版式示例。

    :param slide: 幻灯片对象
    :param palette: 颜色方案
    :return: None
    """
    chart_area = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.7), Inches(7.0), Inches(4.2)
    )
    chart_area.fill.solid()
    chart_area.fill.fore_color.rgb = palette["accent"]
    chart_area.line.fill.background()

    text_box = slide.shapes.add_textbox(Inches(8.2), Inches(1.7), Inches(4.5), Inches(4.2))
    text_frame = text_box.text_frame
    text_frame.text = "洞察一"
    text_frame.add_paragraph().text = "洞察二"


def _add_timeline_layout(slide, palette: dict[str, RGBColor]) -> None:
    """添加时间线页版式示例。

    :param slide: 幻灯片对象
    :param palette: 颜色方案
    :return: None
    """
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(3.2), Inches(11.6), Inches(0.1)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = palette["accent"]
    line.line.fill.background()

    for index in range(4):
        left = Inches(1.2 + index * 2.4)
        node = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, Inches(2.9), Inches(0.6), Inches(0.6))
        node.fill.solid()
        node.fill.fore_color.rgb = palette["accent"]
        node.line.fill.background()

    text_box = slide.shapes.add_textbox(Inches(1.0), Inches(3.6), Inches(11.6), Inches(2.0))
    text_frame = text_box.text_frame
    text_frame.text = "阶段一"
    text_frame.add_paragraph().text = "阶段二"
    text_frame.add_paragraph().text = "阶段三"


def build_templates(output_dir: Path) -> None:
    """生成模板文件。

    :param output_dir: 模板输出目录
    :return: None
    """
    output_dir.mkdir(parents=True, exist_ok=True)

    for spec in TEMPLATE_SPECS:
        prs = Presentation()
        title_slide_layout = prs.slide_layouts[0]
        content_slide_layout = prs.slide_layouts[1]

        cover = prs.slides.add_slide(title_slide_layout)
        _set_slide_background(cover, spec["palette"]["bg"])
        _add_accent_bar(cover, spec["palette"]["accent"], y_inch=0.0, height_inch=0.2)
        if cover.shapes.title:
            cover.shapes.title.text = spec["title"]
            _style_title_shape(cover.shapes.title, spec["palette"]["title"], 44)
        if len(cover.placeholders) > 1:
            cover.placeholders[1].text = spec["subtitle"]

        content = prs.slides.add_slide(content_slide_layout)
        _set_slide_background(content, spec["palette"]["bg"])
        _add_accent_bar(content, spec["palette"]["accent"], y_inch=0.0, height_inch=0.12)
        if content.shapes.title:
            content.shapes.title.text = "内容页标题"
            _style_title_shape(content.shapes.title, spec["palette"]["title"], 32)
        if len(content.placeholders) > 1:
            body = content.placeholders[1].text_frame
            body.text = "要点一"
            body.paragraphs[0].font.color.rgb = spec["palette"]["body"]
            body.paragraphs[0].font.size = Pt(20)

        image_left = prs.slides.add_slide(content_slide_layout)
        _set_slide_background(image_left, spec["palette"]["bg"])
        _add_accent_bar(image_left, spec["palette"]["accent"], y_inch=0.0, height_inch=0.12)
        if image_left.shapes.title:
            image_left.shapes.title.text = "图左文右"
            _style_title_shape(image_left.shapes.title, spec["palette"]["title"], 30)
        _add_image_left_layout(image_left, spec["palette"])

        image_right = prs.slides.add_slide(content_slide_layout)
        _set_slide_background(image_right, spec["palette"]["bg"])
        _add_accent_bar(image_right, spec["palette"]["accent"], y_inch=0.0, height_inch=0.12)
        if image_right.shapes.title:
            image_right.shapes.title.text = "图右文左"
            _style_title_shape(image_right.shapes.title, spec["palette"]["title"], 30)
        _add_image_right_layout(image_right, spec["palette"])

        chart_slide = prs.slides.add_slide(content_slide_layout)
        _set_slide_background(chart_slide, spec["palette"]["bg"])
        _add_accent_bar(chart_slide, spec["palette"]["accent"], y_inch=0.0, height_inch=0.12)
        if chart_slide.shapes.title:
            chart_slide.shapes.title.text = "图表页"
            _style_title_shape(chart_slide.shapes.title, spec["palette"]["title"], 30)
        _add_chart_layout(chart_slide, spec["palette"])

        timeline_slide = prs.slides.add_slide(content_slide_layout)
        _set_slide_background(timeline_slide, spec["palette"]["bg"])
        _add_accent_bar(timeline_slide, spec["palette"]["accent"], y_inch=0.0, height_inch=0.12)
        if timeline_slide.shapes.title:
            timeline_slide.shapes.title.text = "时间线页"
            _style_title_shape(timeline_slide.shapes.title, spec["palette"]["title"], 30)
        _add_timeline_layout(timeline_slide, spec["palette"])

        output_path = output_dir / f"{spec['name']}.pptx"
        _save_presentation(prs, output_path)


def _save_presentation(prs: Presentation, output_path: Path) -> None:
    """安全保存模板文件。

    :param prs: 演示文稿对象
    :param output_path: 输出路径
    :return: None
    """
    try:
        if output_path.exists():
            output_path.unlink()
        prs.save(output_path)
        print(f"生成模板: {output_path}")
    except PermissionError:
        fallback_path = output_path.with_name(f"{output_path.stem}-new{output_path.suffix}")
        prs.save(fallback_path)
        print(f"模板被占用，已生成新文件: {fallback_path}")


def main() -> None:
    """脚本入口。

    :return: None
    """
    output_dir = Path(__file__).resolve().parent.parent / "assets" / "templates"
    build_templates(output_dir)


if __name__ == "__main__":
    main()

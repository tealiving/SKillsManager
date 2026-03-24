#!/usr/bin/env python3
"""通用 PPTX 生成器（Presenton-Lite）。"""

# pip install python-pptx

from __future__ import annotations

from typing import Any

from pptx import Presentation


def build_presentation(slides_data: list[dict[str, Any]], output_path: str = "generated_presentation.pptx") -> None:
    """根据结构化幻灯片数据生成 PPTX 文件。

    :param slides_data: 幻灯片数据列表，每项包含 layout、title、content。
    :param output_path: 输出文件路径。
    :return: None
    """
    prs = Presentation()

    for slide_info in slides_data:
        layout_index = int(slide_info.get("layout", 1))
        slide_layout = prs.slide_layouts[layout_index]
        slide = prs.slides.add_slide(slide_layout)

        title = str(slide_info.get("title", ""))
        if slide.shapes.title:
            slide.shapes.title.text = title

        content = slide_info.get("content", "")
        if len(slide.placeholders) > 1:
            body_shape = slide.placeholders[1]
            if isinstance(content, list):
                text_frame = body_shape.text_frame
                if content:
                    text_frame.text = str(content[0])
                    for item in content[1:]:
                        paragraph = text_frame.add_paragraph()
                        paragraph.text = str(item)
                else:
                    text_frame.text = ""
            else:
                body_shape.text = str(content)

    try:
        prs.save(output_path)
        print(f"成功生成演示文稿: {output_path}")
    except PermissionError:
        print(f"错误: 无法保存文件。请确保 {output_path} 未被其他程序打开。")


def build_demo_slides_data() -> list[dict[str, Any]]:
    """生成示例幻灯片数据用于演示结构。

    :return: 幻灯片数据列表
    """
    # 幻灯片大纲：
    # 1. AI 办公效率提升指南（封面）
    # 2. 为什么需要自动化
    # 3. 三类高频场景
    # 4. 落地步骤与风险控制
    # 5. 总结与行动（结束页）
    return [
        {
            "layout": 0,
            "title": "AI 办公效率提升指南",
            "content": "面向通用职场场景的快速实践\n汇报人：AI Assistant",
        },
        {
            "layout": 1,
            "title": "为什么需要自动化",
            "content": [
                "重复性工作占用大量高价值时间。",
                "信息分散导致沟通成本上升。",
                "标准化流程有助于降低出错率。",
                "AI 可将人从事务性任务中释放出来。",
            ],
        },
        {
            "layout": 1,
            "title": "三类高频场景",
            "content": [
                "文档整理：会议纪要、报告摘要、格式统一。",
                "数据处理：表格清洗、口径对齐、快速可视化。",
                "沟通协作：邮件草拟、任务拆解、进度同步。",
            ],
        },
        {
            "layout": 1,
            "title": "落地步骤与风险控制",
            "content": [
                "选定 1-2 个高收益场景先试点。",
                "建立输入输出模板，减少随意性。",
                "对关键结果设置人工复核。",
                "逐步沉淀可复用的流程与提示词。",
            ],
        },
        {
            "layout": 0,
            "title": "总结与行动",
            "content": "从一个小场景开始，把节省的时间用于更高价值的工作。",
        },
    ]


def main() -> None:
    """生成示例演示文稿文件。

    :return: None
    """
    slides_data = build_demo_slides_data()
    build_presentation(slides_data)


if __name__ == "__main__":
    main()